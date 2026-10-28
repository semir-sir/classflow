import asyncio
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

import database
from lesson_prep import SyllabusDocumentError, extract_syllabus
from lesson_tracker import (
    _outline_review, _topic_selector, handle_callback, handle_document_upload,
    handle_text_input,
)
from telegram.error import NetworkError


TEACHER_ID = 6020889329
TIMEZONE = "Africa/Addis_Ababa"


def make_text_pdf(path):
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
        NameObject("/Encoding"): NameObject("/WinAnsiEncoding"),
    })
    page[NameObject("/Resources")] = DictionaryObject({
        NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})
    })
    content = (
        b"BT /F1 12 Tf 72 720 Td (Chapter 1: Foundations) Tj "
        b"0 -18 Td (1.1 Core concepts) Tj 0 -18 Td (1.1.1 Key terms) Tj ET"
    )
    stream = DecodedStreamObject()
    stream.set_data(content)
    page[NameObject("/Contents")] = writer._add_object(stream)
    writer.add_metadata({"/Title": "Information Systems"})
    with Path(path).open("wb") as output:
        writer.write(output)


def make_outline():
    return {
        "course_title": "Information Systems",
        "chapters": [
            {"title": "1. Foundations", "topics": [
                {"title": "1.1 Core concepts", "children":[
                    {"title": "1.1.1 Key terms", "children": []},
                ]},
                {"title": "1.2 Models", "children": []},
            ]},
            {"title": "2. Analysis", "topics": [
                {"title": "2.1 Requirements", "children": []},
            ]},
        ],
        "warnings": [],
    }


class FakeQuery:
    def __init__(self, data, teacher_id=TEACHER_ID):
        self.data = data
        self.from_user = SimpleNamespace(id=teacher_id)
        self.message = SimpleNamespace(chat_id=123, message_id=456)
        self.text = None
        self.reply_markup = None

    async def answer(self, *args, **kwargs):
        return None

    async def edit_message_text(self, text, reply_markup=None, **kwargs):
        self.text = text
        self.reply_markup = reply_markup


class LessonPrepTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.original_db_path = database.DB_PATH
        database.DB_PATH = Path(self.temp_dir.name) / "lesson-test.sqlite3"
        database.init_db()

    def tearDown(self):
        database.DB_PATH = self.original_db_path
        self.temp_dir.cleanup()

    def stage_syllabus(self, digest="ab" * 32):
        return database.create_pending_syllabus_import(
            TEACHER_ID, "outline.docx", digest, make_outline(),
            "2026-10-28T10:00:00+03:00",
        )["import"]

    def activate_syllabus(self, digest="ab" * 32):
        staged = self.stage_syllabus(digest)
        return database.confirm_syllabus_import(
            staged["id"], TEACHER_ID, "2026-10-28T10:01:00+03:00"
        )

    def make_class(self, date="2026-10-28", session="sad-wed"):
        database.confirm_session(
            session_id=session, day="Wednesday", class_date=date, sections=["A"],
            duration_hours=2, confirmed_at=f"{date}T08:00:00+03:00",
            course="System Analysis and Design", start="04:30", end="06:30",
        )
        return database.get_confirmation(session, date)

    def test_pdf_text_and_docx_heading_extraction(self):
        with tempfile.TemporaryDirectory() as directory:
            pdf_path = Path(directory) / "course.pdf"
            make_text_pdf(pdf_path)
            pdf_outline = extract_syllabus(pdf_path, ".pdf")
            self.assertEqual(pdf_outline["course_title"], "Information Systems")
            self.assertEqual(pdf_outline["chapters"][0]["title"], "Chapter 1: Foundations")
            self.assertEqual(pdf_outline["chapters"][0]["topics"][0]["children"][0]["title"], "1.1.1 Key terms")

            docx_path = Path(directory) / "course.docx"
            document = Document()
            document.core_properties.title = "Information Systems"
            document.add_heading("Chapter 1: Foundations", level=1)
            document.add_heading("1.1 Core concepts", level=2)
            document.add_heading("1.1.1 Key terms", level=3)
            document.add_paragraph("This explanatory paragraph should not become a topic.")
            document.save(docx_path)
            docx_outline = extract_syllabus(docx_path, ".docx")
            self.assertEqual(len(docx_outline["chapters"]), 1)
            self.assertEqual(docx_outline["chapters"][0]["topics"][0]["title"], "1.1 Core concepts")
            self.assertEqual(docx_outline["chapters"][0]["topics"][0]["children"][0]["title"], "1.1.1 Key terms")
            self.assertNotIn("explanatory paragraph", str(docx_outline))

    def test_scanned_empty_and_malformed_pdf_are_reported(self):
        with tempfile.TemporaryDirectory() as directory:
            scanned = Path(directory) / "scan.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=612, height=792)
            with scanned.open("wb") as output:
                writer.write(output)
            with self.assertRaisesRegex(SyllabusDocumentError, "scanned or image-only"):
                extract_syllabus(scanned, ".pdf")
            malformed = Path(directory) / "broken.pdf"
            malformed.write_bytes(b"not a valid pdf")
            with self.assertRaisesRegex(SyllabusDocumentError, "couldn't read this PDF"):
                extract_syllabus(malformed, ".pdf")

            empty_docx = Path(directory) / "empty.docx"
            Document().save(empty_docx)
            with self.assertRaisesRegex(SyllabusDocumentError, "No readable text"):
                extract_syllabus(empty_docx, ".docx")

    def test_pending_review_edit_cancel_duplicate_and_activation(self):
        staged = self.stage_syllabus()
        duplicate = database.create_pending_syllabus_import(
            TEACHER_ID, "copy.docx", "ab" * 32, make_outline(), "2026-10-28T10:02:00+03:00"
        )
        self.assertEqual(duplicate["status"], "already_pending")
        self.assertIsNone(database.get_syllabus(1, TEACHER_ID))
        view = _outline_review(staged["id"], TEACHER_ID, 0)
        self.assertIn("Review Syllabus Import", view[0])
        self.assertTrue(any("Confirm Syllabus" in button.text for row in view[1].inline_keyboard for button in row))
        self.assertTrue(database.cancel_pending_syllabus_import(staged["id"], TEACHER_ID))
        replacement = self.stage_syllabus()
        saved = database.confirm_syllabus_import(replacement["id"], TEACHER_ID, "2026-10-28T10:05:00+03:00")
        self.assertEqual(saved["title"], "Information Systems")
        self.assertEqual(saved["content_sha256"], "ab" * 32)
        duplicate_active = database.create_pending_syllabus_import(
            TEACHER_ID, "copy.docx", "ab" * 32, make_outline(), "2026-10-28T10:06:00+03:00"
        )
        self.assertEqual(duplicate_active["status"], "already_imported")

    def test_syllabus_title_must_be_supplied_before_activation(self):
        outline = make_outline()
        outline["course_title"] = ""
        staged = database.create_pending_syllabus_import(
            TEACHER_ID, "unnamed.pdf", "cd" * 32, outline, "2026-10-28T10:00:00+03:00"
        )["import"]
        with self.assertRaisesRegex(ValueError, "Add a course or syllabus title"):
            database.confirm_syllabus_import(staged["id"], TEACHER_ID, "2026-10-28T10:01:00+03:00")
        self.assertIsNotNone(database.get_pending_syllabus_import(staged["id"], TEACHER_ID))

    def test_existing_class_and_credit_data_survive_additive_migration(self):
        original_path = database.DB_PATH
        database.DB_PATH = Path(self.temp_dir.name) / "existing.sqlite3"
        try:
            with database.connect() as conn:
                conn.executescript("""
                    CREATE TABLE balances (
                        section TEXT PRIMARY KEY CHECK(section IN ('A','B')),
                        hours_remaining REAL NOT NULL CHECK(hours_remaining >= 0)
                    );
                    INSERT INTO balances VALUES ('A', 51.5), ('B', 64.0);
                    CREATE TABLE confirmations (
                        id INTEGER PRIMARY KEY AUTOINCREMENT,
                        session_id TEXT NOT NULL, day TEXT NOT NULL, class_date TEXT NOT NULL,
                        course TEXT NOT NULL, start TEXT NOT NULL, end TEXT NOT NULL,
                        sections TEXT NOT NULL, duration_hours REAL NOT NULL, confirmed_at TEXT NOT NULL,
                        UNIQUE(session_id, class_date)
                    );
                    INSERT INTO confirmations VALUES
                        (7, 'existing', 'Monday', '2026-10-26', 'Existing Course', '08:00', '10:00', 'A', 2, '2026-10-26T08:00:00+03:00');
                """)
            database.init_db()
            self.assertEqual(database.get_balances(), {"A": 51.5, "B": 64.0})
            self.assertEqual(database.get_confirmation("existing", "2026-10-26")["id"], 7)
            self.assertTrue({"syllabus_imports", "syllabuses", "syllabus_chapters", "syllabus_topics",
                             "class_syllabuses", "class_topic_assignments"}.issubset(self.table_names()))
        finally:
            database.DB_PATH = original_path

    def table_names(self):
        with database.connect() as conn:
            return {row[0] for row in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")}

    def test_class_topics_status_notes_unique_progress_and_ownership(self):
        syllabus = self.activate_syllabus()
        class_one = self.make_class()
        self.assertTrue(database.link_syllabus_to_class(class_one["id"], syllabus["id"], TEACHER_ID, "2026-10-28T10:10:00+03:00"))
        self.assertTrue(database.link_syllabus_to_class(class_one["id"], syllabus["id"], TEACHER_ID, "2026-10-28T10:11:00+03:00"))
        topics = database.get_class_syllabus_topics(class_one["id"], syllabus["id"], TEACHER_ID)
        core = next(topic for topic in topics if topic["title"] == "1.1 Core concepts")
        key_terms = next(topic for topic in topics if topic["title"] == "1.1.1 Key terms")
        self.assertTrue(database.toggle_class_topic(class_one["id"], syllabus["id"], core["id"], TEACHER_ID, "2026-10-28T10:12:00+03:00"))
        self.assertFalse(database.toggle_class_topic(class_one["id"], syllabus["id"], core["id"], TEACHER_ID, "2026-10-28T10:13:00+03:00"))
        self.assertTrue(database.set_chapter_topics_selected(
            class_one["id"], syllabus["id"], core["chapter_id"], TEACHER_ID, True, "2026-10-28T10:14:00+03:00"
        ))
        self.assertTrue(database.set_class_topic_status(
            class_one["id"], key_terms["id"], TEACHER_ID, "Partially covered", "2026-10-28T10:15:00+03:00"
        ))
        self.assertTrue(database.set_class_topic_note(
            class_one["id"], key_terms["id"], TEACHER_ID, "Review examples next time.", "2026-10-28T10:16:00+03:00"
        ))
        class_two = self.make_class("2026-11-04", "sad-next-wed")
        database.link_syllabus_to_class(class_two["id"], syllabus["id"], TEACHER_ID, "2026-11-03T10:00:00+03:00")
        self.assertTrue(database.toggle_class_topic(class_two["id"], syllabus["id"], key_terms["id"], TEACHER_ID, "2026-11-04T10:00:00+03:00"))
        database.set_class_topic_status(class_two["id"], key_terms["id"], TEACHER_ID, "Completed", "2026-11-04T12:00:00+03:00")
        progress = database.get_syllabus_progress(syllabus["id"], TEACHER_ID)
        self.assertEqual(progress["total_topics"], 4)
        self.assertEqual(progress["completed_topics"], 1)
        self.assertEqual(progress["remaining_topics"], 3)
        self.assertEqual(progress["percent_complete"], 25.0)
        assignment = next(item for item in database.get_class_topic_assignments(class_one["id"], TEACHER_ID)
                          if item["topic_id"] == key_terms["id"])
        self.assertEqual(assignment["status"], "Partially covered")
        self.assertEqual(assignment["note"], "Review examples next time.")
        self.assertIsNone(database.get_syllabus(syllabus["id"], TEACHER_ID + 1))
        self.assertFalse(database.link_syllabus_to_class(class_one["id"], syllabus["id"], TEACHER_ID + 1, "2026-10-28T12:00:00+03:00"))
        self.assertEqual(database.get_class_syllabus_topics(class_one["id"], syllabus["id"], TEACHER_ID + 1), [])
        self.assertEqual(database.get_balances()["A"], 60.0)

    def test_teacher_callback_review_and_lesson_topic_flow(self):
        staged = self.stage_syllabus()
        context = SimpleNamespace(user_data={})
        confirm_query = FakeQuery(f"syllabus:confirm:{staged['id']}")
        self.assertTrue(asyncio.run(handle_callback(confirm_query, context, TEACHER_ID, TIMEZONE)))
        syllabus = database.list_syllabuses(TEACHER_ID)[0]
        self.assertIn("Syllabus saved", confirm_query.text)
        lesson = self.make_class()
        database.link_syllabus_to_class(lesson["id"], syllabus["id"], TEACHER_ID, "2026-10-28T10:00:00+03:00")

        selector = FakeQuery(f"lesson:topics:{lesson['id']}:{syllabus['id']}:0:0:0")
        asyncio.run(handle_callback(selector, context, TEACHER_ID, TIMEZONE))
        chapter_callback = next(
            button.callback_data
            for row in selector.reply_markup.inline_keyboard for button in row
            if button.callback_data and button.callback_data.startswith("lesson:topics:")
            and button.callback_data.split(":")[4] != "0"
        )
        selector = FakeQuery(chapter_callback)
        asyncio.run(handle_callback(selector, context, TEACHER_ID, TIMEZONE))
        toggle_callback = next(
            button.callback_data
            for row in selector.reply_markup.inline_keyboard for button in row
            if button.callback_data and button.callback_data.startswith("lesson:toggle:")
        )
        toggle_query = FakeQuery(toggle_callback)
        asyncio.run(handle_callback(toggle_query, context, TEACHER_ID, TIMEZONE))
        assigned = database.get_class_topic_assignments(lesson["id"], TEACHER_ID)
        self.assertEqual(len(assigned), 1)

        status_query = FakeQuery(f"lesson:status:{lesson['id']}:{assigned[0]['topic_id']}")
        asyncio.run(handle_callback(status_query, context, TEACHER_ID, TIMEZONE))
        complete_callback = next(
            button.callback_data for row in status_query.reply_markup.inline_keyboard for button in row
            if button.callback_data and button.callback_data.endswith(":c")
        )
        asyncio.run(handle_callback(FakeQuery(complete_callback), context, TEACHER_ID, TIMEZONE))
        self.assertEqual(database.get_class_topic_assignments(lesson["id"], TEACHER_ID)[0]["status"], "Completed")

    def test_topic_assignments_persist_and_cancelled_class_is_not_editable(self):
        syllabus = self.activate_syllabus()
        lesson = self.make_class()
        database.link_syllabus_to_class(lesson["id"], syllabus["id"], TEACHER_ID, "2026-10-28T10:00:00+03:00")
        topic = database.get_class_syllabus_topics(lesson["id"], syllabus["id"], TEACHER_ID)[0]
        self.assertTrue(database.toggle_class_topic(
            lesson["id"], syllabus["id"], topic["id"], TEACHER_ID, "2026-10-28T10:01:00+03:00"
        ))
        database.init_db()
        self.assertEqual(len(database.get_class_topic_assignments(lesson["id"], TEACHER_ID)), 1)
        obligation, created = database.cancel_confirmation_occurrence(
            lesson["id"], "2026-10-29", "2026-10-28T12:00:00+03:00"
        )
        self.assertTrue(created)
        self.assertIsNotNone(obligation)
        self.assertIsNone(database.get_lesson_class(lesson["id"], TEACHER_ID))
        self.assertEqual(database.get_class_topic_assignments(lesson["id"], TEACHER_ID), [])
        self.assertFalse(database.toggle_class_topic(
            lesson["id"], syllabus["id"], topic["id"], TEACHER_ID, "2026-10-28T12:01:00+03:00"
        ))

    def test_review_and_selector_paginate_large_outlines(self):
        outline = {
            "course_title": "Large Course",
            "chapters": [
                {"title": f"Chapter {number}", "topics": [
                    {"title": f"Topic {number}.{topic}", "children": []}
                    for topic in range(1, 11)
                ]}
                for number in range(1, 10)
            ],
            "warnings": [],
        }
        staged = database.create_pending_syllabus_import(
            TEACHER_ID, "large.docx", "12" * 32, outline, "2026-10-28T10:00:00+03:00"
        )["import"]
        review = _outline_review(staged["id"], TEACHER_ID, 0)
        self.assertTrue(any(
            button.callback_data == f"syllabus:review:{staged['id']}:1"
            for row in review[1].inline_keyboard for button in row
        ))
        syllabus = database.confirm_syllabus_import(staged["id"], TEACHER_ID, "2026-10-28T10:01:00+03:00")
        lesson = self.make_class()
        database.link_syllabus_to_class(lesson["id"], syllabus["id"], TEACHER_ID, "2026-10-28T10:02:00+03:00")
        first_chapter = database.get_syllabus_chapters(syllabus["id"], TEACHER_ID)[0]
        markup = _topic_selector(
            lesson["id"], syllabus["id"], TEACHER_ID, 0, first_chapter["id"], 0
        )[1]
        callbacks = [button.callback_data for row in markup.inline_keyboard for button in row if button.callback_data]
        self.assertTrue(any(value.startswith("lesson:topics:") and value.endswith(":0:1") for value in callbacks))
        self.assertTrue(all(len(value.encode("utf-8")) <= 64 for value in callbacks))

    def test_upload_rejects_unsupported_and_non_teacher_files_without_download(self):
        class Message:
            def __init__(self, document):
                self.document = document
                self.replies = []

            async def reply_text(self, text, **kwargs):
                self.replies.append(text)

        class Document:
            file_name = "malware.exe"
            mime_type = "application/octet-stream"
            file_size = 10

        message = Message(Document())
        update = SimpleNamespace(
            effective_message=message, effective_user=SimpleNamespace(id=TEACHER_ID),
            effective_chat=SimpleNamespace(type="private"),
        )
        asyncio.run(handle_document_upload(update, SimpleNamespace(user_data={}), TEACHER_ID, TIMEZONE, 10_000))
        self.assertIn("PDF and DOCX", message.replies[0])
        unauthorized = SimpleNamespace(
            effective_message=Message(Document()), effective_user=SimpleNamespace(id=TEACHER_ID + 1),
            effective_chat=SimpleNamespace(type="private"),
        )
        asyncio.run(handle_document_upload(unauthorized, SimpleNamespace(user_data={}), TEACHER_ID, TIMEZONE, 10_000))
        self.assertIn("authorized teacher", unauthorized.effective_message.replies[0])

    def test_upload_limit_mime_validation_and_temporary_file_cleanup(self):
        class Message:
            def __init__(self, document):
                self.document = document
                self.replies = []

            async def reply_text(self, text, **kwargs):
                self.replies.append(text)

        class DocumentStub:
            file_name = "outline.pdf"
            mime_type = "application/pdf"
            file_size = 3

            async def get_file(self):
                return FileStub()

        downloaded_paths = []

        class FileStub:
            async def download_to_drive(self, custom_path):
                downloaded_paths.append(custom_path)
                custom_path.write_bytes(b"pdf")

        update_message = Message(DocumentStub())
        update = SimpleNamespace(
            effective_message=update_message, effective_user=SimpleNamespace(id=TEACHER_ID),
            effective_chat=SimpleNamespace(type="private"),
        )
        with patch("lesson_tracker.extract_syllabus", return_value=make_outline()):
            asyncio.run(handle_document_upload(update, SimpleNamespace(user_data={}), TEACHER_ID, TIMEZONE, 100))
        pending = database.get_pending_syllabus_import(1, TEACHER_ID)
        self.assertIsNotNone(pending)
        self.assertEqual(len(downloaded_paths), 1)
        self.assertFalse(downloaded_paths[0].exists())

        oversized = DocumentStub()
        oversized.file_size = 101
        too_big_message = Message(oversized)
        too_big_update = SimpleNamespace(
            effective_message=too_big_message, effective_user=SimpleNamespace(id=TEACHER_ID),
            effective_chat=SimpleNamespace(type="private"),
        )
        asyncio.run(handle_document_upload(too_big_update, SimpleNamespace(user_data={}), TEACHER_ID, TIMEZONE, 100))
        self.assertIn("upload limit", too_big_message.replies[0])

        mismatch = DocumentStub()
        mismatch.mime_type = "text/plain"
        mismatch_message = Message(mismatch)
        mismatch_update = SimpleNamespace(
            effective_message=mismatch_message, effective_user=SimpleNamespace(id=TEACHER_ID),
            effective_chat=SimpleNamespace(type="private"),
        )
        asyncio.run(handle_document_upload(mismatch_update, SimpleNamespace(user_data={}), TEACHER_ID, TIMEZONE, 100))
        self.assertIn("does not match", mismatch_message.replies[0])

    def test_telegram_download_failure_is_handled_without_logging_secrets(self):
        class Message:
            document = SimpleNamespace(file_name="outline.pdf", mime_type="application/pdf", file_size=1)
            replies = []

            async def reply_text(self, text, **kwargs):
                self.replies.append(text)

        class BrokenDocument:
            file_name = "outline.pdf"
            mime_type = "application/pdf"
            file_size = 1

            async def get_file(self):
                raise NetworkError("download failed")

        message = Message()
        message.document = BrokenDocument()
        update = SimpleNamespace(
            effective_message=message, effective_user=SimpleNamespace(id=TEACHER_ID),
            effective_chat=SimpleNamespace(type="private"),
        )
        asyncio.run(handle_document_upload(update, SimpleNamespace(user_data={}), TEACHER_ID, TIMEZONE, 100))
        self.assertIn("couldn't download", message.replies[0])

    def test_blank_document_can_be_repaired_manually_before_activation(self):
        outline = {"course_title": "", "chapters": [], "warnings": []}
        staged = database.create_pending_syllabus_import(
            TEACHER_ID, "blank.pdf", "ef" * 32, outline, "2026-10-28T10:00:00+03:00"
        )["import"]
        context = SimpleNamespace(user_data={})
        query = FakeQuery(f"syllabus:edit:{staged['id']}:course:0")
        asyncio.run(handle_callback(query, context, TEACHER_ID, TIMEZONE))
        self.assertEqual(context.user_data["syllabus_edit"]["action"], "course")
        message = SimpleNamespace(
            text="  Database Systems  ",
            reply_text=None,
        )
        async def reply_text(text, **kwargs):
            message.reply_text = text
        message.reply_text = reply_text
        update = SimpleNamespace(
            effective_message=message, effective_user=SimpleNamespace(id=TEACHER_ID),
            effective_chat=SimpleNamespace(type="private"),
        )
        fake_bot = SimpleNamespace(edit_message_text=lambda **kwargs: None)
        # The callback-provided review message is edited through the bot API.
        async def edit_message_text(**kwargs):
            fake_bot.edited = kwargs
        fake_bot.edit_message_text = edit_message_text
        asyncio.run(handle_text_input(update, SimpleNamespace(user_data=context.user_data, bot=fake_bot), TEACHER_ID, TIMEZONE))
        saved = database.get_pending_syllabus_import(staged["id"], TEACHER_ID)
        self.assertEqual(saved["outline"]["course_title"], "Database Systems")


if __name__ == "__main__":
    unittest.main()
